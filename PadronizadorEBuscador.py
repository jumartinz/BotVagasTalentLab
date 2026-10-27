#!/usr/bin/env python3
"""
Analisador de vagas + coletor automático.

Uso:
    python analisador_vagas.py                      # coleta e gera vagas.xlsx
    python analisador_vagas.py --so-com-cluster     # descarta vagas sem cluster
    python analisador_vagas.py --termos "estágio dados" "devops" --limite 100
    python analisador_vagas.py --fontes gupy        # só uma fonte
    python analisador_vagas.py --manual             # modo antigo (colar texto)
    python analisador_vagas.py --posts posts.txt    # onde salvar os posts prontos
                                                    # (padrão: <saida>_posts.txt)

Dependências:  pip install requests openpyxl
"""

import argparse
import html
import os
import re
import sys
import time
from datetime import datetime

# ---------------------------------------------------------------------------
# 1. CRITÉRIOS DE CLASSIFICAÇÃO (os seus, com \b nos termos curtos para evitar
#    falsos positivos: "cto" casava com "director", "java" com "javascript",
#    "aws" com "laws", "cio" com "negócio" etc.)
# ---------------------------------------------------------------------------
CRITERIOS = {
    "tipo_contratacao": {
        "Estágio": [r"\bestágio\b", r"\bestagio\b", r"\bintern\b", r"\binternship\b"],
        "Trainee": [r"\btrainee\b"],
        "Jovem Aprendiz": [r"\bjovem aprendiz\b", r"\baprendiz\b", r"\bapprentice\b"],
        "Emprego (CLT)": [r"\bclt\b", r"\befetivo\b"],
        "Pessoa Jurídica (PJ)": [r"\bpj\b", r"\bpessoa jurídica\b", r"\bpessoa juridica\b", r"\bprestador de serviço\b"],
        "Temporário": [r"\btemporário\b", r"\btemporario\b", r"\btemporary\b"],
        "Meio período": [r"\bmeio período\b", r"\bmeio periodo\b", r"\bpart[- ]time\b"],
        "Voluntariado": [r"\bvolunt[aá]ri[ao]\b", r"\bvolunteer\b"],
    },
    "modelo_trabalho": {
        "Remoto": [r"\bremoto\b", r"\bremote\b", r"\bhome office\b", r"\b100% remoto\b"],
        "Híbrido": [r"\bhíbrido\b", r"\bhibrido\b", r"\bhybrid\b"],
        "No local": [r"\bno local\b", r"\bpresencial\b", r"\bon-site\b", r"\bonsite\b"],
    },
    "nivel_vaga": {
        "Early Career": [r"\bestágio\b", r"\bestagio\b", r"\btrainee\b", r"\bjúnior\b", r"\bjunior\b", r"\bjr\b", r"\bentry[- ]level\b"],
        "Professional/Experienced": [r"\bpleno\b", r"\bpl\b", r"\bsênior\b", r"\bsenior\b", r"\bsr\b", r"\bmid[- ]level\b"],
        "Advanced/Expert": [r"\bstaff\b", r"\bprincipal\b", r"\bespecialista\b", r"\bspecialist\b", r"\bexpert\b"],
        "Leadership": [r"\bmanager\b", r"\bgerente\b", r"\bgestor\b", r"\bhead\b", r"\bdiretor\b", r"\bdirector\b", r"\bcoordenador\b", r"\blead\b"],
    },
    "clusters": {
        "Cyber": [r"seguran[çc]a cibern[eé]tica", r"seguran[çc]a da informa[çc][aã]o", r"cyber\s*security", r"defensive security", r"offensive security", r"\bappsec\b", r"\binfosec\b"],
        "IA": [r"intelig[eê]ncia artificial", r"\bia\b", r"\bai\b", r"machine learning", r"agentes de ia", r"ai scientist", r"multi[- ]agents?"],
        "Dados": [r"data science", r"ci[eê]ncia de dados", r"data analytics", r"engenharia de dados", r"data engineering", r"business intelligence", r"\bbi\b"],
        "Dev": [r"an[aá]lise e desenvolvimento de sistemas", r"engenharia de software", r"software architecture", r"\bjava\b", r"\.net", r"full\s*stack", r"backend", r"desenvolvedor", r"developer"],
        "Cloud & Infraestrutura": [r"computa[çc][aã]o em nuvem", r"\bcloud\b", r"\bdevops\b", r"\baws\b", r"\bazure\b", r"\bgcp\b", r"infraestrutura"],
        "Creative Tech": [r"design digital", r"web design", r"game development", r"front\s*end", r"ui/ux", r"ux design"],
        "Robotics & Systems": [r"mecatr[oô]nica", r"engenharia de computa[çc][aã]o", r"sistemas embarcados", r"rob[oó]tica"],
        "Tech Business & Liderança": [r"gest[aã]o de ti", r"tech management", r"product manager", r"gest[aã]o estrat[eé]gica", r"project management", r"\bcto\b", r"\bcio\b", r"scrum master"],
    },
}

PADROES_AFIRMATIVA = [
    r"\bpcd\b",
    r"pessoa[s]? com defici[eê]ncia",
    r"\bmulher(es)?\b",
    r"\bpreto[s]?\b",
    r"\bpreta[s]?\b",
    r"\bpardo[s]?\b",
    r"\bparda[s]?\b",
    r"\bnegro[s]?\b",
    r"\bnegra[s]?\b",
    r"\bafirmativ[ao]\b",
    r"diversidade",
    r"inclus[aã]o",
]

# Termos de busca padrão (derivados dos seus clusters)
TERMOS_PT = [
    "estágio tecnologia", "desenvolvedor", "dados", "segurança da informação",
    "inteligência artificial", "cloud", "devops", "ux design", "robótica",
    "gestão de ti",
]
TERMOS_EN = ["developer", "data", "security", "machine learning", "devops", "cloud", "intern"]


# ---------------------------------------------------------------------------
# 2. CLASSIFICAÇÃO
# ---------------------------------------------------------------------------
def _casa(padroes, texto):
    return any(re.search(p, texto) for p in padroes)


def verificar_afirmativa(texto: str) -> bool:
    return _casa(PADROES_AFIRMATIVA, texto.lower())


def extrair_titulo(texto: str, eh_afirmativa: bool) -> str:
    linhas = [l.strip() for l in texto.strip().splitlines() if l.strip()]
    titulo_base = linhas[0] if linhas else "Vaga sem título informado"
    titulo_base = re.sub(r"^(t[ií]tulo da vaga|t[ií]tulo|vaga|cargo):\s*", "", titulo_base, flags=re.IGNORECASE)
    return _marcar_afirmativa(titulo_base, eh_afirmativa)


def _marcar_afirmativa(titulo: str, eh_afirmativa: bool) -> str:
    if eh_afirmativa and not re.search(r"afirmativ[ao]", titulo, re.IGNORECASE):
        return f"{titulo} - Vaga Afirmativa"
    return titulo


def gerar_resumo(texto: str, limite: int = 2000) -> str:
    limpo = re.sub(r"\s+", " ", texto).strip()
    return limpo if len(limpo) <= limite else limpo[: limite - 3] + "..."


def classificar_vaga(texto: str, titulo: str | None = None, dicas: dict | None = None) -> dict:
    """
    texto  : descrição completa da vaga
    titulo : título já conhecido (vindo da API). Se None, usa a 1ª linha do texto.
    dicas  : categorias já informadas pela fonte, ex.
             {"tipo": {"Estágio"}, "modelo": {"Remoto"}}
    """
    dicas = dicas or {}
    texto_lower = texto.lower()
    eh_afirmativa = verificar_afirmativa(f"{titulo or ''} {texto}")

    if titulo:
        titulo_final = _marcar_afirmativa(titulo.strip(), eh_afirmativa)
    else:
        titulo_final = extrair_titulo(texto, eh_afirmativa)

    res = {
        "Título": titulo_final,
        "Tipo de contratação": set(dicas.get("tipo", set())),
        "Modelo de trabalho": set(dicas.get("modelo", set())),
        "Nível da vaga": set(),
        "Clusters identificados": set(),
        "Vaga afirmativa": eh_afirmativa,
        "Resumo": gerar_resumo(texto),
    }

    for cat, pads in CRITERIOS["tipo_contratacao"].items():
        if _casa(pads, texto_lower):
            res["Tipo de contratação"].add(cat)
    for cat, pads in CRITERIOS["modelo_trabalho"].items():
        if _casa(pads, texto_lower):
            res["Modelo de trabalho"].add(cat)

    # Nível: olha primeiro o título (descrições longas citam "sênior", "lead"
    # etc. em outros contextos). Só cai para o texto todo se o título não ajudar.
    primeira_linha = texto.strip().splitlines()[0] if texto.strip() else ""
    titulo_lower = (titulo or primeira_linha).lower()
    for fonte_nivel in (titulo_lower, texto_lower):
        for nivel, pads in CRITERIOS["nivel_vaga"].items():
            if _casa(pads, fonte_nivel):
                res["Nível da vaga"].add(nivel)
        if res["Nível da vaga"]:
            break

    for cluster, pads in CRITERIOS["clusters"].items():
        if _casa(pads, texto_lower):
            res["Clusters identificados"].add(cluster)

    return res


# ---------------------------------------------------------------------------
# 3. COLETORES (APIs/feeds públicos — sem scraping de HTML)
# ---------------------------------------------------------------------------
def _sessao():
    import requests
    s = requests.Session()
    s.headers.update({"User-Agent": "Mozilla/5.0 (analisador-de-vagas; uso pessoal)", "Accept": "application/json"})
    return s


def _limpar_html(txt: str) -> str:
    txt = re.sub(r"<(br|/p|/li|/h\d)\s*/?>", "\n", txt or "", flags=re.I)
    txt = re.sub(r"<[^>]+>", " ", txt)
    return html.unescape(txt)


def coletar_gupy(termos, limite, pausa=1.0):
    """Gupy: maior plataforma de vagas do Brasil (endpoint público do portal)."""
    url = "https://portal.api.gupy.io/api/v1/jobs"
    tipos = {
        "vacancy_type_effective": "Emprego (CLT)",
        "vacancy_type_internship": "Estágio",
        "vacancy_type_temporary": "Temporário",
        "vacancy_type_apprentice": "Jovem Aprendiz",
    }
    modelos = {"remote": "Remoto", "hybrid": "Híbrido", "on-site": "No local"}
    s, vagas = _sessao(), []

    for termo in termos:
        offset, total = 0, 0
        while total < limite:
            lote = min(10, limite - total)  # a API pagina de 10 em 10
            try:
                r = s.get(url, params={"jobName": termo, "limit": lote, "offset": offset}, timeout=20)
                r.raise_for_status()
                dados = r.json().get("data", [])
            except Exception as e:
                print(f"  [gupy] falha em '{termo}': {e}", file=sys.stderr)
                break
            if not dados:
                break
            for j in dados:
                dicas = {"tipo": set(), "modelo": set()}
                if j.get("type") in tipos:
                    dicas["tipo"].add(tipos[j["type"]])
                if j.get("workplaceType") in modelos:
                    dicas["modelo"].add(modelos[j["workplaceType"]])
                elif j.get("isRemoteWork"):
                    dicas["modelo"].add("Remoto")
                local = " - ".join(x for x in (j.get("city"), j.get("state")) if x) or j.get("country", "")
                vagas.append({
                    "titulo": j.get("title") or j.get("name", ""),
                    "empresa": j.get("companyName") or j.get("careerPageName", ""),
                    "fonte": "Gupy",
                    "local": local,
                    "data": (j.get("publishedDate") or "")[:10],
                    "link": j.get("applicationUrl") or j.get("jobUrl") or j.get("applyUrl") or j.get("careerPageUrl", ""),
                    "descricao": _limpar_html(j.get("description", "")),
                    "dicas": dicas,
                })
            total += len(dados)
            offset += len(dados)
            if len(dados) < lote:
                break
            time.sleep(pausa)
        print(f"  [gupy] '{termo}': {total} vagas")
        time.sleep(pausa)
    return vagas


def coletar_remotive(termos, limite, pausa=1.0):
    """Remotive: vagas remotas internacionais (API pública)."""
    s, vagas = _sessao(), []
    for termo in termos:
        try:
            r = s.get("https://remotive.com/api/remote-jobs", params={"search": termo, "limit": limite}, timeout=20)
            r.raise_for_status()
            dados = r.json().get("jobs", [])
        except Exception as e:
            print(f"  [remotive] falha em '{termo}': {e}", file=sys.stderr)
            continue
        for j in dados:
            dicas = {"tipo": set(), "modelo": {"Remoto"}}
            if j.get("job_type") == "internship":
                dicas["tipo"].add("Estágio")
            elif j.get("job_type") == "part_time":
                dicas["tipo"].add("Meio período")
            vagas.append({
                "titulo": j.get("title", ""),
                "empresa": j.get("company_name", ""),
                "fonte": "Remotive",
                "local": j.get("candidate_required_location", ""),
                "data": (j.get("publication_date") or "")[:10],
                "link": j.get("url", ""),
                "descricao": _limpar_html(j.get("description", "")),
                "dicas": dicas,
            })
        print(f"  [remotive] '{termo}': {len(dados)} vagas")
        time.sleep(pausa)
    return vagas


def coletar_remoteok(termos, limite, pausa=1.0):
    """RemoteOK: lista única; o filtro por termo é feito aqui no cliente."""
    s, vagas = _sessao(), []
    try:
        r = s.get("https://remoteok.com/api", timeout=20)
        r.raise_for_status()
        dados = [d for d in r.json() if isinstance(d, dict) and d.get("position")]
    except Exception as e:
        print(f"  [remoteok] falha: {e}", file=sys.stderr)
        return vagas
    termos_l = [t.lower() for t in termos]
    for j in dados:
        base = f"{j.get('position','')} {' '.join(j.get('tags') or [])}".lower()
        if termos_l and not any(t in base for t in termos_l):
            continue
        vagas.append({
            "titulo": j.get("position", ""),
            "empresa": j.get("company", ""),
            "fonte": "RemoteOK",
            "local": j.get("location", "") or "Remoto",
            "data": (j.get("date") or "")[:10],
            "link": j.get("url", ""),
            "descricao": _limpar_html(j.get("description", "")),
            "dicas": {"tipo": set(), "modelo": {"Remoto"}},
        })
        if len(vagas) >= limite * max(len(termos_l), 1):
            break
    print(f"  [remoteok]: {len(vagas)} vagas")
    return vagas


FONTES = {
    "gupy": (coletar_gupy, "pt"),
    "remotive": (coletar_remotive, "en"),
    "remoteok": (coletar_remoteok, "en"),
}


# ---------------------------------------------------------------------------
# 3.1 FORMATAÇÃO DE POSTS (texto simples + emojis: funciona em WhatsApp,
#     Telegram, LinkedIn e Instagram sem depender de markup)
# ---------------------------------------------------------------------------
HASHTAGS_CLUSTER = {
    "Cyber": "#SegurancaDaInformacao",
    "IA": "#InteligenciaArtificial",
    "Dados": "#Dados",
    "Dev": "#Desenvolvimento",
    "Cloud & Infraestrutura": "#Cloud",
    "Creative Tech": "#UXDesign",
    "Robotics & Systems": "#Robotica",
    "Tech Business & Liderança": "#GestaoDeTI",
}
SEPARADOR_POSTS = "\n" + "─" * 30 + "\n\n"


def resumir_para_post(descricao: str, limite: int = 350) -> str:
    """Trecho curto da descrição, cortado em fim de frase ou de palavra."""
    limpo = re.sub(r"\s+", " ", descricao or "").strip()
    if len(limpo) <= limite:
        return limpo
    corte = limpo[:limite]
    ponto = max(corte.rfind(". "), corte.rfind("! "), corte.rfind("? "))
    if ponto > limite * 0.5:
        return corte[: ponto + 1]
    return corte.rsplit(" ", 1)[0].rstrip(",;:-") + "…"


def formatar_post(v: dict) -> str:
    """Monta o texto pronto para postar. Linhas sem informação são omitidas."""
    def juntar(chave):
        return ", ".join(sorted(v[chave]))

    linhas = [f"💼 {v['Título']}"]
    if v.get("empresa"):
        linhas.append(f"🏢 {v['empresa'].strip()}")

    local_modelo = " · ".join(x for x in (v.get("local"), juntar("Modelo de trabalho")) if x)
    if local_modelo:
        linhas.append(f"📍 {local_modelo}")
    if v["Tipo de contratação"]:
        linhas.append(f"📄 {juntar('Tipo de contratação')}")
    if v["Nível da vaga"]:
        linhas.append(f"🎯 {juntar('Nível da vaga')}")
    if v["Clusters identificados"]:
        linhas.append(f"🧩 Área: {juntar('Clusters identificados')}")
    if v["Vaga afirmativa"]:
        linhas.append("💜 Vaga afirmativa")

    resumo = resumir_para_post(v.get("descricao", ""))
    if resumo:
        linhas += ["", resumo]

    if v.get("link"):
        linhas += ["", f"🔗 Candidate-se: {v['link']}"]
    if v.get("fonte"):
        linhas.append(f"Fonte: {v['fonte']}")

    tags = ["#Vagas", "#TalentLab"] + [HASHTAGS_CLUSTER[c] for c in sorted(v["Clusters identificados"])]
    if v["Vaga afirmativa"]:
        tags.append("#VagaAfirmativa")
    linhas += ["", " ".join(tags)]
    return "\n".join(linhas)


def salvar_posts(registros: list[dict], caminho: str):
    with open(caminho, "w", encoding="utf-8") as f:
        f.write(SEPARADOR_POSTS.join(v["Post"] for v in registros) + "\n")


# ---------------------------------------------------------------------------
# 4. PLANILHA
# ---------------------------------------------------------------------------
COLUNAS = [
    ("Título", 48), ("Empresa", 26), ("Fonte", 11), ("Local", 24),
    ("Tipo de contratação", 22), ("Modelo de trabalho", 18), ("Nível da vaga", 24),
    ("Clusters identificados", 30), ("Vaga afirmativa", 14), ("Data de publicação", 16),
    ("Link", 14), ("Resumo", 80), ("Post formatado", 60),
]


def gerar_planilha(registros: list[dict], caminho: str, termos_usados: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    cab_font = Font(name=F, bold=True, color="FFFFFF", size=10)
    cab_fill = PatternFill("solid", fgColor="1F3864")
    corpo = Font(name=F, size=10)
    link_font = Font(name=F, size=10, color="0563C1", underline="single")
    fino = Side(style="thin", color="D9D9D9")
    borda = Border(bottom=fino)

    wb = Workbook()
    ws = wb.active
    ws.title = "Vagas"

    for c, (nome, larg) in enumerate(COLUNAS, 1):
        cel = ws.cell(row=1, column=c, value=nome)
        cel.font, cel.fill = cab_font, cab_fill
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = larg
    ws.row_dimensions[1].height = 28

    for i, v in enumerate(registros, 2):
        linha = [
            v["Título"], v["empresa"], v["fonte"], v["local"],
            ", ".join(sorted(v["Tipo de contratação"])) or "Não identificado",
            ", ".join(sorted(v["Modelo de trabalho"])) or "Não identificado",
            ", ".join(sorted(v["Nível da vaga"])) or "Não identificado",
            ", ".join(sorted(v["Clusters identificados"])) or "Não identificado",
            "Sim" if v["Vaga afirmativa"] else "Não",
            v["data"], v["link"], v["Resumo"], v["Post"],
        ]
        for c, val in enumerate(linha, 1):
            cel = ws.cell(row=i, column=c, value=val)
            cel.font, cel.border = corpo, borda
            cel.alignment = Alignment(vertical="top", wrap_text=(c in (1, 5, 6, 7, 8)))
        if v["link"]:
            cel = ws.cell(row=i, column=11, value="Abrir vaga")
            cel.hyperlink = v["link"]
            cel.font = link_font

    n = max(len(registros) + 1, 2)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{n}"

    # ----- Aba de resumo (tudo em fórmulas, recalcula se você editar a aba Vagas)
    rs = wb.create_sheet("Resumo")
    rs.column_dimensions["A"].width = 34
    rs.column_dimensions["B"].width = 12
    negrito = Font(name=F, bold=True, size=10)
    sec_fill = PatternFill("solid", fgColor="D9E1F2")

    rs["A1"] = "Resumo das vagas coletadas"
    rs["A1"].font = Font(name=F, bold=True, size=14)
    rs["A2"] = f"Gerado em {datetime.now():%d/%m/%Y %H:%M}" + (f" | termos: {termos_usados}" if termos_usados else "")
    rs["A2"].font = Font(name=F, italic=True, size=9, color="595959")

    rs["A4"], rs["B4"] = "Total de vagas", f"=COUNTA(Vagas!A2:A{n})"
    rs["A5"], rs["B5"] = "Vagas afirmativas", f'=COUNTIF(Vagas!I2:I{n},"Sim")'
    for r in (4, 5):
        rs[f"A{r}"].font, rs[f"B{r}"].font = negrito, corpo

    def bloco(titulo, col, categorias, linha_ini, contem=True):
        rs.cell(row=linha_ini, column=1, value=titulo).font = negrito
        rs.cell(row=linha_ini, column=2, value="Vagas").font = negrito
        rs.cell(row=linha_ini, column=1).fill = sec_fill
        rs.cell(row=linha_ini, column=2).fill = sec_fill
        r = linha_ini + 1
        for cat in categorias:
            rs.cell(row=r, column=1, value=cat).font = corpo
            rng = f"Vagas!${col}$2:${col}${n}"
            if contem:  # célula com várias categorias separadas por ", "
                f = f'=SUMPRODUCT(--ISNUMBER(SEARCH("|"&A{r}&"|","|"&SUBSTITUTE({rng},", ","|")&"|")))'
            else:
                f = f"=COUNTIF({rng},A{r})"
            rs.cell(row=r, column=2, value=f).font = corpo
            r += 1
        return r + 1

    prox = 7
    prox = bloco("Por cluster", "H", list(CRITERIOS["clusters"]), prox)
    prox = bloco("Por tipo de contratação", "E", list(CRITERIOS["tipo_contratacao"]), prox)
    prox = bloco("Por modelo de trabalho", "F", list(CRITERIOS["modelo_trabalho"]), prox)
    prox = bloco("Por nível", "G", list(CRITERIOS["nivel_vaga"]), prox)
    fontes = sorted({v["fonte"] for v in registros}) or ["Gupy"]
    bloco("Por fonte", "C", fontes, prox, contem=False)

    wb.save(caminho)


def deduplicar(vagas):
    vistos, saida = set(), []
    for v in vagas:
        chave = v["link"] or (v["titulo"].lower(), v["empresa"].lower())
        if chave in vistos:
            continue
        vistos.add(chave)
        saida.append(v)
    return saida


def coletar_e_classificar(fontes, termos_pt, termos_en, limite, so_com_cluster=False):
    brutas = []
    for nome in fontes:
        func, idioma = FONTES[nome]
        print(f"Coletando de {nome}...")
        brutas += func(termos_pt if idioma == "pt" else termos_en, limite)

    brutas = deduplicar(brutas)
    registros = []
    for v in brutas:
        res = classificar_vaga(f"{v['titulo']}\n{v['descricao']}", titulo=v["titulo"], dicas=v["dicas"])
        res.update({k: v[k] for k in ("empresa", "fonte", "local", "data", "link", "descricao")})
        if so_com_cluster and not res["Clusters identificados"]:
            continue
        res["Post"] = formatar_post(res)
        registros.append(res)
    return registros


def modo_manual():
    print("Cole a descrição da vaga (Ctrl+D no Linux/Mac, Ctrl+Z + Enter no Windows):")
    texto = sys.stdin.read()
    if not texto.strip():
        print("Nenhum texto informado.")
        return
    r = classificar_vaga(texto)
    print("\n" + "=" * 60)
    for k in ("Título", "Tipo de contratação", "Modelo de trabalho", "Nível da vaga", "Clusters identificados"):
        v = r[k]
        print(f"{k:24}: {', '.join(sorted(v)) if isinstance(v, set) else v or 'Não identificado'}")
    print("-" * 60)
    print(r["Resumo"])
    print("=" * 60)


def main():
    ap = argparse.ArgumentParser(description="Coleta vagas e gera planilha.")
    ap.add_argument("--manual", action="store_true", help="modo antigo: colar uma descrição")
    ap.add_argument("--fontes", nargs="+", choices=list(FONTES), default=list(FONTES))
    ap.add_argument("--termos", nargs="+", default=TERMOS_PT, help="termos de busca em português (Gupy)")
    ap.add_argument("--termos-en", nargs="+", default=TERMOS_EN, help="termos em inglês (fontes remotas)")
    ap.add_argument("--limite", type=int, default=50, help="máx. de vagas por termo e fonte")
    ap.add_argument("--saida", default="vagas.xlsx")
    ap.add_argument("--posts", default=None, help="arquivo .txt com os posts prontos (padrão: <saida>_posts.txt)")
    ap.add_argument("--so-com-cluster", action="store_true", help="descarta vagas sem nenhum cluster")
    a = ap.parse_args()

    if a.manual:
        return modo_manual()

    registros = coletar_e_classificar(a.fontes, a.termos, a.termos_en, a.limite, a.so_com_cluster)
    if not registros:
        print("Nenhuma vaga coletada. Verifique a conexão ou os termos de busca.")
        return
    gerar_planilha(registros, a.saida, termos_usados=", ".join(a.termos))
    caminho_posts = a.posts or os.path.splitext(a.saida)[0] + "_posts.txt"
    salvar_posts(registros, caminho_posts)
    print(f"\n{len(registros)} vagas salvas em {a.saida}")
    print(f"{len(registros)} posts prontos em {caminho_posts}")


if __name__ == "__main__":
    main()