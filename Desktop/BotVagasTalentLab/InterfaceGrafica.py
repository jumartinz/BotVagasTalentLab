#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FIAP · Coletor de Vagas de Tecnologia
=====================================

Busca vagas de tecnologia em fontes públicas brasileiras, aplica as regras de
negócio da FIAP e exporta tudo em lote para Excel.

Fontes
------
- Gupy Portal      : API pública de busca de vagas (JSON), sem chave.
- GitHub Issues    : repositórios comunitários brasileiros de vagas
                     (frontendbr/vagas, backend-br/vagas, ...). Sem chave
                     (token opcional para ampliar o rate limit).
- Remotive         : API pública de vagas remotas (filtra Brasil/LATAM/Worldwide).
- Jooble           : API oficial (requer chave gratuita em jooble.org/api/about).

Uso
---
    pip install -r requirements.txt
    python fiap_vagas_tech.py              # abre a interface gráfica
    python fiap_vagas_tech.py --selftest   # teste offline das regras + Excel

Estrutura do código
-------------------
    Constantes / normalização  -> utilitários de texto
    Regras de negócio          -> classify_*, detect_affirmative, process_job
    Fontes                     -> JobSource e subclasses (requests / BeautifulSoup)
    JobHunter                  -> orquestra a coleta (thread, parada, progresso)
    export_to_excel            -> pandas + openpyxl
    App                        -> GUI CustomTkinter
"""
from __future__ import annotations

import html
import queue
import re
import sys
import threading
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Callable, Optional

import pandas as pd
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:  # GUI é opcional para o --selftest
    import customtkinter as ctk
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:  # pragma: no cover
    ctk = None

# --------------------------------------------------------------------------- #
# Constantes de negócio (FIAP)
# --------------------------------------------------------------------------- #
CONTRATACAO = ["Emprego (CLT)", "Meio Período", "Estágio", "Temporário",
               "Trainee", "Pessoa Jurídica (PJ)", "Voluntariado"]
MODALIDADES_TRABALHO = ["No local (presencial)", "Híbrido", "Remoto"]
CLUSTERS = ["Cyber", "IA", "Dados", "Dev", "Cloud & Infraestrutura",
            "Creative Tech", "Robotic & Systems", "Tech Business e Liderança"]
NIVEIS = ["Early Career", "Professional / Experienced",
          "Advanced / Expert", "Leadership"]

MAX_TITLE = 58
MAX_DESC = 2000
DEFAULT_VALIDITY_DAYS = 30
AFFIRMATIVE_SUFFIX = " - Afirmativa (minoritária)"

COLUMNS = [
    "Modalidade de Contratação",
    "Título da Oportunidade",
    "Descrição da Oportunidade",
    "Modalidade de Trabalho",
    "Cluster e Blocos",
    "Nível da Vaga",
    "Local da Oportunidade",
    "Data de Publicação",
    "Data de Expiração",
    "Link de Inscrição",
    "Geração de pacote de CVs ao expirar",
    "Aprovado (Uso Administrativo)",
    "Emprego em destaque para docente",
]
FIXED_PACOTE_CV = "Não"
FIXED_APROVADO = "SIM"
FIXED_DESTAQUE_DOCENTE = "Não"

GITHUB_REPOS = [
    "frontendbr/vagas", "backend-br/vagas", "androiddevbr/vagas",
    "soujava/vagas-java", "phpdevbr/vagas", "datascience-br/vagas",
    "qa-brasil/vagas",
]
DEFAULT_TERMS = [
    "desenvolvedor", "engenheiro de dados", "cloud", "devops",
    "segurança da informação", "inteligência artificial",
]


# --------------------------------------------------------------------------- #
# Utilitários de texto
# --------------------------------------------------------------------------- #
def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s or "")
                   if not unicodedata.combining(c))


def norm(s: str) -> str:
    """minúsculas + sem acentos (para casar palavras-chave)."""
    return strip_accents(s or "").lower()


def kw_regex(kw: str) -> re.Pattern:
    return re.compile(r"(?<![a-z0-9])" + re.escape(kw) + r"(?![a-z0-9])")


_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def sanitize(value):
    return _ILLEGAL_XML.sub("", value) if isinstance(value, str) else value


def truncate(text: str, limit: int) -> str:
    """Corta em `limit` caracteres (respeitando palavras) com reticências."""
    text = text or ""
    if len(text) <= limit:
        return text
    cut = text[: limit - 1]
    sp = cut.rfind(" ")
    if sp > limit * 0.6:
        cut = cut[:sp]
    return cut.rstrip(" ,;:-–·") + "…"


def html_to_text(raw: str) -> str:
    """HTML/Markdown -> texto limpo, preservando quebras de linha e bullets."""
    if not raw:
        return ""
    text = raw
    if "<" in raw and ">" in raw:
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup.find_all(["script", "style"]):
            tag.decompose()
        for br in soup.find_all("br"):
            br.replace_with("\n")
        for li in soup.find_all("li"):
            li.insert(0, "• ")
        for tag in soup.find_all(["p", "div", "li", "ul", "ol", "tr",
                                  "h1", "h2", "h3", "h4", "h5", "h6"]):
            tag.append("\n")
        text = soup.get_text("")
    text = html.unescape(text)
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)          # imagens md
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)      # links md
    text = re.sub(r"(?m)^\s{0,3}#{1,6}\s*", "", text)         # headers md
    text = re.sub(r"(?m)^\s*[-*]\s+", "• ", text)             # bullets md
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def parse_date(value) -> Optional[date]:
    """Aceita ISO (qualquer variação) ou dd/mm/aaaa. Retorna None se falhar."""
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value)
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    try:
        if m:
            return date(int(m[1]), int(m[2]), int(m[3]))
        m = re.search(r"(\d{2})/(\d{2})/(\d{4})", s)
        if m:
            return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None
    return None


# --------------------------------------------------------------------------- #
# Modelo de dados
# --------------------------------------------------------------------------- #
@dataclass
class RawJob:
    """Vaga bruta, como veio da fonte (antes das regras da FIAP)."""
    source: str
    title: str
    url: str
    description: str = ""
    company: str = ""
    location: str = ""
    published: Optional[str] = None
    deadline: Optional[str] = None
    contract_hint: Optional[str] = None   # tipo informado pela fonte
    workmode_hint: Optional[str] = None   # remote / hybrid / on-site
    labels: str = ""                      # labels/tags (GitHub, etc.)
    salary_hint: str = ""
    affirmative_flag: bool = False        # ex.: Gupy "disabilities"


# --------------------------------------------------------------------------- #
# Regras de negócio
# --------------------------------------------------------------------------- #
# (peso, [palavras-chave normalizadas]) — "Dev" tem peso menor por ser genérico.
_CLUSTER_DEFS: dict[str, tuple[int, list[str]]] = {
    "Cyber": (2, [
        "seguranca da informacao", "ciberseguranca", "seguranca cibernetica",
        "cyber", "pentest", "pentester", "soc", "siem", "appsec", "devsecops",
        "infosec", "vulnerabilidade", "vulnerabilidades", "red team",
        "blue team", "firewall", "owasp", "forense digital", "iso 27001",
        "threat hunting", "incident response", "resposta a incidentes",
        "seguranca ofensiva", "iam", "gestao de identidades"]),
    "IA": (2, [
        "inteligencia artificial", "machine learning", "aprendizado de maquina",
        "ia generativa", "genai", "llm", "llms", "deep learning", "nlp",
        "visao computacional", "mlops", "redes neurais", "pytorch",
        "tensorflow", "prompt engineer", "engenheiro de ia", "ia", "ai",
        "ai engineer", "ml engineer", "openai", "rag"]),
    "Dados": (2, [
        "engenheiro de dados", "analista de dados", "cientista de dados",
        "engenharia de dados", "ciencia de dados", "analise de dados",
        "data engineer", "data scientist", "data analyst", "data science",
        "power bi", "etl", "sql", "big data", "data warehouse", "analytics",
        "tableau", "spark", "airflow", "databricks", "dbt", "dba",
        "banco de dados", "business intelligence", "bi", "looker",
        "data lake", "data lakehouse", "governanca de dados"]),
    "Dev": (1, [
        "desenvolvedor", "desenvolvedora", "developer", "programador",
        "software engineer", "engenheiro de software", "backend", "back-end",
        "back end", "frontend", "front-end", "front end", "fullstack",
        "full stack", "full-stack", "java", "python", "javascript",
        "typescript", "react", "angular", "vue", "node", "nodejs", ".net",
        "c#", "c++", "php", "golang", "ruby", "kotlin", "swift", "android",
        "ios", "mobile", "flutter", "react native", "qa", "sdet",
        "testes automatizados", "qualidade de software", "api", "apis",
        "spring", "laravel", "django", "flask", "scala", "rust", "dev",
        "analista de sistemas", "analista desenvolvedor", "software"]),
    "Cloud & Infraestrutura": (2, [
        "cloud", "aws", "azure", "gcp", "google cloud", "devops", "sre",
        "kubernetes", "k8s", "docker", "terraform", "infraestrutura",
        "infra", "redes", "network", "linux", "sysadmin", "ansible",
        "administrador de sistemas", "help desk", "service desk",
        "suporte tecnico", "ci/cd", "observabilidade", "platform engineer",
        "engenharia de plataforma", "cloud engineer", "noc",
        "administrador de redes", "virtualizacao", "vmware"]),
    "Creative Tech": (2, [
        "ux", "ui", "ux/ui", "ui/ux", "product designer", "design de produto",
        "designer", "game", "games", "jogos", "unity", "unreal",
        "realidade virtual", "realidade aumentada", "vr", "xr", "3d",
        "motion", "motion design", "creative technologist", "figma",
        "webgl", "animacao", "game designer", "ux writer", "ux research",
        "design system"]),
    "Robotic & Systems": (2, [
        "robotica", "robotico", "iot", "embarcado", "embarcados", "embedded",
        "firmware", "automacao industrial", "automacao", "plc", "arduino",
        "ros", "drone", "drones", "mecatronica", "scada", "eletronica",
        "hardware", "microcontrolador", "microcontroladores", "sistemas ciberfisicos",
        "raspberry pi", "industria 4.0", "visao robotica"]),
    "Tech Business e Liderança": (2, [
        "product manager", "gerente de produto", "product owner",
        "scrum master", "agile coach", "agilista", "gerente de projetos",
        "gerente de projeto", "project manager", "analista de negocios",
        "business analyst", "pmo", "cto", "tech lead", "lider tecnico",
        "gerente de ti", "coordenador de ti", "diretor de tecnologia",
        "transformacao digital", "customer success", "pre-vendas",
        "pre vendas", "presales", "consultor de tecnologia", "gestao de ti",
        "itil", "engineering manager", "gerente de engenharia",
        "gerente de tecnologia", "head de tecnologia", "head of engineering",
        "product ops", "analista de produto", "growth"]),
}
CLUSTER_PATTERNS = {
    name: (weight, [kw_regex(k) for k in kws])
    for name, (weight, kws) in _CLUSTER_DEFS.items()
}
MIN_CLUSTER_SCORE = 2


def classify_cluster(title: str, text: str) -> Optional[str]:
    """Pontua cada cluster (título pesa 3x, descrição até 3 ocorrências)."""
    nt, nb = norm(title), norm(text)
    best, best_score = None, 0
    for name, (weight, patterns) in CLUSTER_PATTERNS.items():
        score = 0
        for p in patterns:
            if p.search(nt):
                score += 3 * weight
            score += min(len(p.findall(nb)), 3) * weight
        if score > best_score:
            best, best_score = name, score
    return best if best_score >= MIN_CLUSTER_SCORE else None


_GUPY_TYPE_MAP = {
    "vacancy_type_effective": "Emprego (CLT)",
    "vacancy_type_temporary": "Temporário",
    "vacancy_type_internship": "Estágio",
    "vacancy_type_apprentice": "Estágio",
    "vacancy_type_trainee": "Trainee",
}


def _normalize_contract_hint(hint: Optional[str]) -> Optional[str]:
    if not hint:
        return None
    if hint in CONTRATACAO:
        return hint
    h = norm(hint)
    if h in _GUPY_TYPE_MAP:
        return _GUPY_TYPE_MAP[h]
    if re.search(r"volunt", h):
        return "Voluntariado"
    if re.search(r"intern|estag", h):
        return "Estágio"
    if re.search(r"trainee", h):
        return "Trainee"
    if re.search(r"temporar|temporary", h):
        return "Temporário"
    if re.search(r"part[ _-]?time|meio[ _]periodo", h):
        return "Meio Período"
    if re.search(r"contract|freelanc|autonom|\bpj\b|outsourc", h):
        return "Pessoa Jurídica (PJ)"
    if re.search(r"effective|full[ _-]?time|clt|efetivo", h):
        return "Emprego (CLT)"
    return None


def classify_contract(title: str, labels: str, body: str,
                      hint: Optional[str] = None) -> str:
    """Valida/infere a modalidade de contratação (sempre um valor da lista)."""
    head = norm(f"{title} {labels}")
    if re.search(r"(?<![a-z])trainee", head):
        return "Trainee"
    if re.search(r"estagi|internship|(?<![a-z])intern(?![a-z])|aprendiz", head):
        return "Estágio"
    if re.search(r"volunt", head):
        return "Voluntariado"
    if re.search(r"temporari|temporary", head):
        return "Temporário"
    hinted = _normalize_contract_hint(hint)
    if hinted:
        return hinted
    if re.search(r"meio periodo|part[- ]?time", head):
        return "Meio Período"
    if re.search(r"(?<![a-z])pj(?![a-z])|pessoa juridica|cnpj", head):
        return "Pessoa Jurídica (PJ)"
    if re.search(r"(?<![a-z])clt(?![a-z])", head):
        return "Emprego (CLT)"
    nb = norm(body)
    if re.search(r"vaga volunt|trabalho volunt", nb):
        return "Voluntariado"
    if re.search(r"meio periodo|part[- ]?time", nb):
        return "Meio Período"
    if re.search(r"contrato temporario|vaga temporaria|prazo determinado", nb):
        return "Temporário"
    has_clt = bool(re.search(r"(?<![a-z])clt(?![a-z])", nb))
    has_pj = bool(re.search(r"(?<![a-z])pj(?![a-z])|pessoa juridica|cnpj", nb))
    if has_pj and not has_clt:
        return "Pessoa Jurídica (PJ)"
    return "Emprego (CLT)"


def classify_workmode(title_ctx: str, body: str,
                      hint: Optional[str] = None) -> str:
    """Retorna estritamente: No local (presencial) | Híbrido | Remoto."""
    if hint:
        h = norm(hint)
        if "hybrid" in h or "hibrid" in h:
            return "Híbrido"
        if "remot" in h:
            return "Remoto"
        if "on-site" in h or "onsite" in h or "presencial" in h:
            return "No local (presencial)"
    for chunk in (norm(title_ctx), norm(body)):
        if re.search(r"hibrid|hybrid", chunk):
            return "Híbrido"
        if re.search(r"remot|home[- ]?office|anywhere|trabalho a distancia|100% online",
                     chunk):
            return "Remoto"
    return "No local (presencial)"


_LEAD = re.compile(r"(?<![a-z])(gerente|coordenador|coordenadora|head|diretor|"
                   r"diretora|lider|supervisor|supervisora|cto|cio|vp|tech lead|"
                   r"team lead|engineering manager)(?![a-z])")
_EARLY = re.compile(r"(?<![a-z])(junior|jr|estagi\w*|trainee|aprendiz|assistente|"
                    r"auxiliar|entry level|iniciante)(?![a-z])")
_ADV = re.compile(r"(?<![a-z])(senior|sr|especialista|principal|staff|arquiteto|"
                  r"arquiteta|architect|expert)(?![a-z])")
_PRO = re.compile(r"(?<![a-z])(pleno|pl|mid|middle|analista)(?![a-z])")


def classify_level(title: str, labels: str, body: str, contract: str) -> str:
    if contract in ("Estágio", "Trainee"):
        return "Early Career"
    head = norm(f"{title} {labels}")
    if _LEAD.search(head):
        return "Leadership"
    if _EARLY.search(head):
        return "Early Career"
    if _ADV.search(head):
        return "Advanced / Expert"
    if _PRO.search(head):
        return "Professional / Experienced"
    nb = norm(body)
    counts = {
        "Early Career": len(re.findall(r"(?<![a-z])(junior|jr)(?![a-z])", nb)),
        "Advanced / Expert": len(re.findall(r"(?<![a-z])(senior|sr|especialista)(?![a-z])", nb)),
        "Professional / Experienced": len(re.findall(r"(?<![a-z])pleno(?![a-z])", nb)),
    }
    level, hits = max(counts.items(), key=lambda kv: kv[1])
    return level if hits >= 1 else "Professional / Experienced"


_AFFIRM_PATTERNS: dict[str, list[re.Pattern]] = {
    "PCD": [kw_regex("pcd"), kw_regex("pcds"), kw_regex("pwd"),
            re.compile(r"pessoas? com deficiencia"), re.compile(r"deficiente")],
    "Mulheres": [re.compile(r"(?<![a-z])mulher(?:es)?(?![a-z])"),
                 re.compile(r"(?<![a-z])women(?![a-z])")],
    "Pessoas Pretas": [re.compile(r"pessoas? (?:pretas?|negras?)"),
                       re.compile(r"(?<![a-z])negros?(?![a-z])"),
                       re.compile(r"afrodescendentes?")],
    "Pardas": [re.compile(r"pessoas? pardas?"), re.compile(r"(?<![a-z])pardos?(?![a-z])")],
    "Indígenas": [re.compile(r"indigenas?")],
    "LGBT": [re.compile(r"lgbt\w*")],
}
_AFFIRM_CONTEXT = re.compile(
    r"afirmativ|exclusiv|prioridade|prioritari|diversidade|inclusa|inclusiv|"
    r"incentiv|estimul|encorajamos|destinad|dedicad|reservad|somente|apenas|"
    r"vagas? (?:para|aberta)|affirmative|diversity|inclusion")
_AFFIRM_EXPLICIT = re.compile(r"vagas?\s+afirmativ|acao afirmativa|affirmative action")


def detect_affirmative(text: str, flag: bool = False) -> list[str]:
    """Grupos minoritários citados em contexto afirmativo na descrição."""
    nt = norm(text)
    explicit = bool(_AFFIRM_EXPLICIT.search(nt))
    groups: list[str] = ["PCD"] if flag else []
    for group, patterns in _AFFIRM_PATTERNS.items():
        if group in groups:
            continue
        for p in patterns:
            hit = False
            for m in p.finditer(nt):
                window = nt[max(0, m.start() - 90): m.end() + 90]
                if explicit or _AFFIRM_CONTEXT.search(window):
                    hit = True
                    break
            if hit:
                groups.append(group)
                break
    return groups


def format_title(base_title: str, affirmative: bool) -> str:
    """Título ≤ 58 caracteres; afirmativas recebem o sufixo obrigatório."""
    base = re.sub(r"\s+", " ", base_title or "").strip(" -–|:")
    base = re.sub(r"\s*[-–|]\s*vagas?\s+afirmativa.*$", "", base, flags=re.I)
    if affirmative:
        room = MAX_TITLE - len(AFFIRMATIVE_SUFFIX)
        return truncate(base, room).rstrip("… ") + AFFIRMATIVE_SUFFIX \
            if len(base) > room else base + AFFIRMATIVE_SUFFIX
    return truncate(base, MAX_TITLE)


_SALARY = re.compile(
    r"R\$\s?\d[\d\.\,]*(?:\s*(?:a|à|até|-|–)\s*(?:R\$\s?)?\d[\d\.\,]*)?"
    r"(?:\s*(?:/|por|ao)\s*(?:mês|mes|hora|ano))?", re.I)
_WORKLOAD = re.compile(
    r"(\d{1,2})\s?(?:h|horas?)\s?(?:semanais|por semana|/\s?semana|diarias|diárias|por dia)"
    r"|(?:carga hor[áa]ria[^\d]{0,15})(\d{1,2})\s?(?:h|horas)", re.I)


def extract_salary(text: str, hint: str = "") -> str:
    if hint and re.search(r"\d", hint):
        return truncate(re.sub(r"\s+", " ", hint).strip(), 80)
    m = _SALARY.search(text or "")
    return truncate(m.group(0).strip(), 80) if m else ""


def extract_workload(text: str) -> str:
    m = _WORKLOAD.search(text or "")
    if not m:
        return ""
    raw = m.group(0)
    if re.search(r"di[aá]rias|por dia", raw, re.I):
        return f"{m.group(1)}h diárias"
    hours = m.group(1) or m.group(2)
    return f"{hours}h semanais"


def build_description(contract: str, mode: str, location: str, workload: str,
                      salary: str, company: str, body: str) -> str:
    """Cabeçalho estruturado + corpo, garantindo ≤ 2000 caracteres no total."""
    lines = []
    if company:
        lines.append(f"Empresa: {truncate(company, 80)}")
    lines += [f"Modalidade de contratação: {contract}",
              f"Modalidade de trabalho: {mode}",
              f"Local: {truncate(location, 100)}"]
    if workload:
        lines.append(f"Carga horária: {workload}")
    if salary:
        lines.append(f"Salário: {salary}")
    header = "\n".join(lines)
    room = MAX_DESC - len(header) - 2
    body_part = truncate(body, room) if room > 60 else ""
    full = header + (f"\n\n{body_part}" if body_part else "")
    return full[:MAX_DESC]


def process_job(raw: RawJob) -> Optional[dict]:
    """Aplica TODAS as regras da FIAP. Retorna None se não for vaga de tecnologia."""
    body = html_to_text(raw.description)
    title_ctx = f"{raw.title} {raw.labels}"

    cluster = classify_cluster(raw.title, f"{body} {raw.labels}")
    if cluster is None:
        return None

    contract = classify_contract(raw.title, raw.labels, body, raw.contract_hint)
    mode = classify_workmode(title_ctx, body, raw.workmode_hint)
    level = classify_level(raw.title, raw.labels, body, contract)
    groups = detect_affirmative(f"{raw.title} {raw.labels} {body}", raw.affirmative_flag)
    title = format_title(raw.title, bool(groups))

    location = (raw.location or "").strip()
    if not location:
        location = "Remoto (Brasil)" if mode == "Remoto" else "Brasil"

    workload = extract_workload(body)
    salary = extract_salary(body, raw.salary_hint)
    description = build_description(contract, mode, location, workload, salary,
                                    raw.company, body)

    today = date.today()
    published = parse_date(raw.published) or today
    expiry = parse_date(raw.deadline)
    if expiry is None or expiry < published:
        base = max(published, today)
        expiry = base + timedelta(days=DEFAULT_VALIDITY_DAYS)

    return {
        "Modalidade de Contratação": contract,
        "Título da Oportunidade": title,
        "Descrição da Oportunidade": description,
        "Modalidade de Trabalho": mode,
        "Cluster e Blocos": cluster,
        "Nível da Vaga": level,
        "Local da Oportunidade": location,
        "Data de Publicação": published,
        "Data de Expiração": expiry,
        "Link de Inscrição": raw.url,
        "Geração de pacote de CVs ao expirar": FIXED_PACOTE_CV,
        "Aprovado (Uso Administrativo)": FIXED_APROVADO,
        "Emprego em destaque para docente": FIXED_DESTAQUE_DOCENTE,
    }


def validate_row(row: dict) -> list[str]:
    """Checagem final das regras; devolve lista de problemas (vazia = ok)."""
    issues = []
    if row["Modalidade de Contratação"] not in CONTRATACAO:
        issues.append("contratação inválida")
    if len(row["Título da Oportunidade"]) > MAX_TITLE:
        issues.append(f"título > {MAX_TITLE}")
    if len(row["Descrição da Oportunidade"]) > MAX_DESC:
        issues.append(f"descrição > {MAX_DESC}")
    if row["Modalidade de Trabalho"] not in MODALIDADES_TRABALHO:
        issues.append("modalidade de trabalho inválida")
    if row["Cluster e Blocos"] not in CLUSTERS:
        issues.append("cluster inválido")
    if row["Nível da Vaga"] not in NIVEIS:
        issues.append("nível inválido")
    if not str(row["Link de Inscrição"]).startswith("http"):
        issues.append("link inválido")
    if row["Data de Expiração"] < row["Data de Publicação"]:
        issues.append("expiração < publicação")
    if (row["Geração de pacote de CVs ao expirar"], row["Aprovado (Uso Administrativo)"],
            row["Emprego em destaque para docente"]) != (FIXED_PACOTE_CV, FIXED_APROVADO,
                                                         FIXED_DESTAQUE_DOCENTE):
        issues.append("campos fixos alterados")
    return issues


# --------------------------------------------------------------------------- #
# Fontes de dados
# --------------------------------------------------------------------------- #
class SourceError(Exception):
    """Falha recuperável ao consultar uma fonte (rede, HTTP, JSON)."""


def build_session() -> requests.Session:
    s = requests.Session()
    retry = Retry(total=3, backoff_factor=1.0,
                  status_forcelist=(429, 500, 502, 503, 504),
                  allowed_methods=("GET", "POST"), respect_retry_after_header=True)
    adapter = HTTPAdapter(max_retries=retry)
    s.mount("https://", adapter)
    s.mount("http://", adapter)
    s.headers.update({
        "User-Agent": "FIAP-VagasTech/1.0 (+coletor academico; contato: carreiras@fiap.com.br)",
        "Accept": "application/json",
    })
    return s


class JobSource:
    name = "base"
    max_pages = 10          # teto de páginas por consulta
    timeout = 20

    def __init__(self, session: requests.Session,
                 log: Callable[[str, str], None]):
        self.session = session
        self.log = log

    def queries(self, terms: list[str]) -> list[str]:
        return list(terms)

    def search(self, query: str, page: int) -> list[RawJob]:  # pragma: no cover
        raise NotImplementedError

    def _json(self, method: str, url: str, **kwargs):
        try:
            resp = self.session.request(method, url, timeout=self.timeout, **kwargs)
        except requests.exceptions.Timeout as exc:
            raise SourceError(f"timeout em {url}") from exc
        except requests.exceptions.ConnectionError as exc:
            raise SourceError(f"falha de conexão ({url})") from exc
        except requests.exceptions.RequestException as exc:
            raise SourceError(f"erro de requisição: {exc}") from exc
        if resp.status_code == 404:
            raise SourceError(f"404 não encontrado: {url}")
        if resp.status_code in (401, 403):
            raise SourceError(f"acesso negado/limite ({resp.status_code}) em {url}")
        if not resp.ok:
            raise SourceError(f"HTTP {resp.status_code} em {url}")
        try:
            return resp.json()
        except ValueError as exc:
            raise SourceError("resposta não é JSON válido") from exc


class GupySource(JobSource):
    name = "Gupy"
    URL = "https://portal.api.gupy.io/api/job"
    PAGE = 20
    max_pages = 15

    def search(self, query, page):
        data = self._json("GET", self.URL, params={
            "name": query, "limit": self.PAGE, "offset": page * self.PAGE})
        items = (data or {}).get("data") or []
        jobs = []
        for it in items:
            try:
                url = it.get("jobUrl") or ""
                if not url:
                    continue
                country = norm(it.get("country") or "")
                if country and country not in ("brasil", "brazil", "br"):
                    continue
                city = it.get("city") or it.get("addressCity") or ""
                state = it.get("state") or it.get("addressState") or ""
                wt = (it.get("workplaceType") or "").lower()
                if not wt and it.get("isRemoteWork"):
                    wt = "remote"
                jobs.append(RawJob(
                    source=self.name, title=it.get("name") or "", url=url,
                    description=it.get("description") or "",
                    company=it.get("careerPageName") or "",
                    location=" - ".join(p for p in (city, state) if p),
                    published=it.get("publishedDate"),
                    deadline=it.get("applicationDeadline"),
                    contract_hint=it.get("type"), workmode_hint=wt,
                    affirmative_flag=bool(it.get("disabilities"))))
            except Exception as exc:  # item malformado não derruba a página
                self.log(f"Gupy: item ignorado ({exc})", "warn")
        return jobs


class GitHubIssuesSource(JobSource):
    name = "GitHub Vagas"
    PAGE = 40
    max_pages = 5

    def __init__(self, session, log, token: str = "", repos=None):
        super().__init__(session, log)
        self.token = token.strip()
        self.repos = repos or GITHUB_REPOS

    def queries(self, terms):
        return list(self.repos)

    def search(self, query, page):
        headers = {"Accept": "application/vnd.github+json"}
        if self.token:
            headers["Authorization"] = f"Bearer {self.token}"
        data = self._json("GET", f"https://api.github.com/repos/{query}/issues",
                          headers=headers,
                          params={"state": "open", "per_page": self.PAGE,
                                  "page": page + 1})
        jobs = []
        for it in data or []:
            try:
                if "pull_request" in it or not it.get("html_url"):
                    continue
                raw_title = it.get("title") or ""
                tags = re.findall(r"\[([^\]]{2,40})\]", raw_title)
                clean_title = re.sub(r"\[[^\]]*\]", "", raw_title).strip(" -–|")
                label_names = [l.get("name", "") for l in it.get("labels", [])]
                known = re.compile(r"remot|clt|pj|junior|pleno|senior|estagi|home|"
                                   r"hibrid|presencial|trainee|full|part", re.I)
                loc = next((t for t in tags if not known.search(t)), "")
                jobs.append(RawJob(
                    source=self.name, title=clean_title or raw_title,
                    url=it["html_url"], description=it.get("body") or "",
                    location=loc, published=it.get("created_at"),
                    labels=" ".join(label_names + tags)))
            except Exception as exc:
                self.log(f"GitHub: item ignorado ({exc})", "warn")
        return jobs


class RemotiveSource(JobSource):
    name = "Remotive"
    URL = "https://remotive.com/api/remote-jobs"
    CATEGORIES = ["software-dev", "data", "devops", "qa", "product", "design"]
    max_pages = 1
    _OK_LOC = re.compile(r"brazil|brasil|latam|latin america|south america|"
                         r"americas|worldwide|anywhere")

    def queries(self, terms):
        return list(self.CATEGORIES)

    def search(self, query, page):
        data = self._json("GET", self.URL, params={"category": query, "limit": 100})
        jobs = []
        for it in (data or {}).get("jobs", []):
            try:
                loc = it.get("candidate_required_location") or ""
                if not self._OK_LOC.search(norm(loc)):
                    continue
                jobs.append(RawJob(
                    source=self.name, title=it.get("title") or "",
                    url=it.get("url") or "", description=it.get("description") or "",
                    company=it.get("company_name") or "",
                    location=f"Remoto - {loc}" if loc else "Remoto (Brasil)",
                    published=it.get("publication_date"),
                    contract_hint=it.get("job_type"), workmode_hint="remote",
                    salary_hint=it.get("salary") or ""))
            except Exception as exc:
                self.log(f"Remotive: item ignorado ({exc})", "warn")
        return [j for j in jobs if j.url]


class JoobleSource(JobSource):
    name = "Jooble"
    max_pages = 10

    def __init__(self, session, log, api_key: str = ""):
        super().__init__(session, log)
        self.api_key = api_key.strip()

    def search(self, query, page):
        if not self.api_key:
            raise SourceError("chave da API Jooble não informada")
        data = self._json("POST", f"https://br.jooble.org/api/{self.api_key}",
                          json={"keywords": query, "location": "Brasil",
                                "page": str(page + 1)})
        jobs = []
        for it in (data or {}).get("jobs", []):
            try:
                if not it.get("link"):
                    continue
                jobs.append(RawJob(
                    source=self.name, title=it.get("title") or "", url=it["link"],
                    description=it.get("snippet") or "", company=it.get("company") or "",
                    location=it.get("location") or "", published=it.get("updated"),
                    contract_hint=it.get("type"), salary_hint=it.get("salary") or ""))
            except Exception as exc:
                self.log(f"Jooble: item ignorado ({exc})", "warn")
        return jobs


# --------------------------------------------------------------------------- #
# Orquestração da coleta
# --------------------------------------------------------------------------- #
class JobHunter:
    """Executa a coleta; pensado para rodar em thread separada."""

    def __init__(self, sources: list[JobSource], terms: list[str], max_pages: int,
                 stop_event: threading.Event,
                 on_log: Callable[[str, str], None],
                 on_progress: Callable[[int, int, str], None],
                 on_job: Callable[[dict], None], delay: float = 0.7):
        self.sources, self.terms, self.max_pages = sources, terms, max_pages
        self.stop, self.log = stop_event, on_log
        self.on_progress, self.on_job, self.delay = on_progress, on_job, delay
        self.rows: list[dict] = []
        self.seen: set[str] = set()
        self.discarded = 0

    def run(self) -> list[dict]:
        plan = []
        for src in self.sources:
            qs = src.queries(self.terms)
            plan.append((src, qs, max(1, min(self.max_pages, src.max_pages))))
        total = sum(len(q) * p for _, q, p in plan) or 1
        base = 0
        for src, qs, pages in plan:
            if self.stop.is_set():
                break
            self.log(f"▶ Fonte {src.name}: {len(qs)} consulta(s) × até {pages} pág.", "info")
            consecutive_errors = 0
            for qi, query in enumerate(qs):
                if self.stop.is_set() or consecutive_errors >= 3:
                    break
                for page in range(pages):
                    if self.stop.is_set():
                        break
                    self.on_progress(base + qi * pages + page, total,
                                     f"{src.name} · {query} · pág. {page + 1}")
                    try:
                        raws = src.search(query, page)
                        consecutive_errors = 0
                    except SourceError as exc:
                        consecutive_errors += 1
                        self.log(f"{src.name} [{query}]: {exc}", "warn")
                        break
                    except Exception as exc:  # nunca derrubar o app
                        consecutive_errors += 1
                        self.log(f"{src.name} [{query}]: erro inesperado: {exc!r}", "error")
                        break
                    added = self._ingest(raws)
                    self.log(f"{src.name} [{query}] pág. {page + 1}: "
                             f"{len(raws)} lidas, {added} novas", "info")
                    if not raws:
                        break
                    self.stop.wait(self.delay)
            if consecutive_errors >= 3:
                self.log(f"{src.name}: muitas falhas seguidas — fonte ignorada.", "error")
            base += len(qs) * pages
            self.on_progress(base, total, f"{src.name} concluída")
        return self.rows

    def _ingest(self, raws: list[RawJob]) -> int:
        added = 0
        for raw in raws:
            key = (raw.url or "").split("?")[0].strip().lower()
            if not key or key in self.seen:
                continue
            self.seen.add(key)
            try:
                row = process_job(raw)
            except Exception as exc:
                self.log(f"Vaga ignorada ({raw.title[:40]}): {exc!r}", "warn")
                continue
            if row is None:
                self.discarded += 1
                continue
            problems = validate_row(row)
            if problems:
                self.log(f"Vaga fora das regras ({', '.join(problems)}): "
                         f"{row['Título da Oportunidade']}", "warn")
                continue
            self.rows.append(row)
            self.on_job(row)
            added += 1
        return added


# --------------------------------------------------------------------------- #
# Exportação para Excel
# --------------------------------------------------------------------------- #
def export_to_excel(rows: list[dict], path: str) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    clean = [{c: sanitize(r.get(c)) for c in COLUMNS} for r in rows]
    df = pd.DataFrame(clean, columns=COLUMNS)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Vagas", index=False)
        ws = writer.sheets["Vagas"]
        n = len(df)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        head_fill = PatternFill("solid", fgColor="ED145B")  # rosa FIAP
        widths = [24, 40, 70, 22, 24, 26, 26, 16, 16, 45, 22, 22, 22]

        for ci, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w
            c = ws.cell(row=1, column=ci)
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[1].height = 34

        date_cols = {COLUMNS.index("Data de Publicação") + 1,
                     COLUMNS.index("Data de Expiração") + 1}
        link_col = COLUMNS.index("Link de Inscrição") + 1
        desc_col = COLUMNS.index("Descrição da Oportunidade") + 1
        for r in range(2, n + 2):
            ws.row_dimensions[r].height = 90
            for ci in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=r, column=ci)
                cell.font = Font(name="Arial", size=10)
                cell.border = border
                cell.alignment = Alignment(vertical="top",
                                           wrap_text=ci in (2, desc_col, 7))
                if ci in date_cols:
                    cell.number_format = "DD/MM/YYYY"
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                if ci == link_col and cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        last = max(n + 1, 500)  # validação também para linhas novas digitadas
        lists = {
            "Modalidade de Contratação": CONTRATACAO,
            "Modalidade de Trabalho": MODALIDADES_TRABALHO,
            "Cluster e Blocos": CLUSTERS,
            "Nível da Vaga": NIVEIS,
            "Geração de pacote de CVs ao expirar": ["Sim", "Não"],
            "Aprovado (Uso Administrativo)": ["SIM", "Não"],
            "Emprego em destaque para docente": ["Sim", "Não"],
        }
        for col_name, options in lists.items():
            letter = get_column_letter(COLUMNS.index(col_name) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                                allow_blank=True, showErrorMessage=True)
            dv.error = "Valor fora da lista permitida pela FIAP."
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{last}")


# --------------------------------------------------------------------------- #
# Interface gráfica (CustomTkinter)
# --------------------------------------------------------------------------- #
FIAP_PINK = "#ED145B"
FIAP_PINK_HOVER = "#C40F4B"

_Base = ctk.CTk if ctk is not None else object


class App(_Base):
    TREE_COLS = (("titulo", "Título", 320), ("contratacao", "Contratação", 130),
                 ("trabalho", "Trabalho", 130), ("cluster", "Cluster", 170),
                 ("nivel", "Nível", 170), ("local", "Local", 160))

    def __init__(self):
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")
        super().__init__()
        self.title("FIAP · Coletor de Vagas de Tecnologia")
        self.geometry("1220x780")
        self.minsize(1000, 660)

        self.events: queue.Queue = queue.Queue()
        self.stop_event = threading.Event()
        self.worker: Optional[threading.Thread] = None
        self.rows: list[dict] = []

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self._build_sidebar()
        self._build_main()
        self._apply_tree_style()
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.after(100, self._poll)
        self._log("Pronto. Configure as fontes e clique em “Iniciar Busca”.", "info")

    # ---- construção da UI -------------------------------------------------
    def _build_sidebar(self):
        sb = ctk.CTkScrollableFrame(self, width=290, corner_radius=0, label_text="")
        sb.grid(row=0, column=0, sticky="nsew")
        ctk.CTkLabel(sb, text="FIAP", font=ctk.CTkFont(size=30, weight="bold"),
                     text_color=FIAP_PINK).pack(anchor="w", padx=14, pady=(14, 0))
        ctk.CTkLabel(sb, text="Coletor de Vagas Tech",
                     font=ctk.CTkFont(size=14)).pack(anchor="w", padx=14, pady=(0, 14))

        ctk.CTkLabel(sb, text="Termos de busca (um por linha)",
                     font=ctk.CTkFont(weight="bold")).pack(anchor="w", padx=14)
        self.terms_box = ctk.CTkTextbox(sb, height=130)
        self.terms_box.pack(fill="x", padx=14, pady=(4, 12))
        self.terms_box.insert("1.0", "\n".join(DEFAULT_TERMS))

        ctk.CTkLabel(sb, text="Fontes",
                     font=ctk.CTkFont(weight="bold")).pack(anchor="w", padx=14)
        self.var_gupy = tk.BooleanVar(value=True)
        self.var_github = tk.BooleanVar(value=True)
        self.var_remotive = tk.BooleanVar(value=False)
        self.var_jooble = tk.BooleanVar(value=False)
        for text, var in (("Gupy (API pública)", self.var_gupy),
                          ("GitHub · repositórios de vagas BR", self.var_github),
                          ("Remotive (remoto, filtra Brasil/LATAM)", self.var_remotive),
                          ("Jooble (requer chave)", self.var_jooble)):
            ctk.CTkCheckBox(sb, text=text, variable=var,
                            fg_color=FIAP_PINK, hover_color=FIAP_PINK_HOVER
                            ).pack(anchor="w", padx=14, pady=3)

        self.pages_label = ctk.CTkLabel(sb, text="Páginas por consulta: 3")
        self.pages_label.pack(anchor="w", padx=14, pady=(12, 0))
        self.pages_slider = ctk.CTkSlider(
            sb, from_=1, to=10, number_of_steps=9, button_color=FIAP_PINK,
            button_hover_color=FIAP_PINK_HOVER, progress_color=FIAP_PINK,
            command=lambda v: self.pages_label.configure(
                text=f"Páginas por consulta: {int(v)}"))
        self.pages_slider.set(3)
        self.pages_slider.pack(fill="x", padx=14, pady=(2, 10))

        ctk.CTkLabel(sb, text="Chave API Jooble (opcional)").pack(anchor="w", padx=14)
        self.jooble_key = ctk.CTkEntry(sb, show="•", placeholder_text="xxxxxxxx-xxxx-...")
        self.jooble_key.pack(fill="x", padx=14, pady=(2, 8))
        ctk.CTkLabel(sb, text="Token GitHub (opcional, ↑ rate limit)").pack(anchor="w", padx=14)
        self.gh_token = ctk.CTkEntry(sb, show="•", placeholder_text="ghp_...")
        self.gh_token.pack(fill="x", padx=14, pady=(2, 12))

        ctk.CTkLabel(sb, text="Aparência").pack(anchor="w", padx=14)
        ctk.CTkOptionMenu(sb, values=["Escuro", "Claro", "Sistema"],
                          fg_color=FIAP_PINK, button_color=FIAP_PINK_HOVER,
                          command=self._set_appearance).pack(fill="x", padx=14, pady=(2, 14))

    def _build_main(self):
        main = ctk.CTkFrame(self, fg_color="transparent")
        main.grid(row=0, column=1, sticky="nsew", padx=16, pady=16)
        main.grid_columnconfigure(0, weight=1)
        main.grid_rowconfigure(3, weight=1)

        bar = ctk.CTkFrame(main, fg_color="transparent")
        bar.grid(row=0, column=0, sticky="ew")
        self.btn_start = ctk.CTkButton(bar, text="▶  Iniciar Busca", height=40,
                                       fg_color=FIAP_PINK, hover_color=FIAP_PINK_HOVER,
                                       command=self.start_search)
        self.btn_stop = ctk.CTkButton(bar, text="■  Parar", height=40, state="disabled",
                                      fg_color="#6b7280", hover_color="#4b5563",
                                      command=self.stop_search)
        self.btn_export = ctk.CTkButton(bar, text="⬇  Exportar para Excel", height=40,
                                        state="disabled", fg_color="#16a34a",
                                        hover_color="#15803d", command=self.export_excel)
        for i, b in enumerate((self.btn_start, self.btn_stop, self.btn_export)):
            b.grid(row=0, column=i, padx=(0, 10), sticky="w")

        prog = ctk.CTkFrame(main, fg_color="transparent")
        prog.grid(row=1, column=0, sticky="ew", pady=(14, 4))
        prog.grid_columnconfigure(0, weight=1)
        self.progress = ctk.CTkProgressBar(prog, progress_color=FIAP_PINK, height=14)
        self.progress.set(0)
        self.progress.grid(row=0, column=0, sticky="ew")
        self.pct_label = ctk.CTkLabel(prog, text="0%", width=48)
        self.pct_label.grid(row=0, column=1, padx=(10, 0))
        self.status_label = ctk.CTkLabel(main, text="Aguardando…", anchor="w")
        self.status_label.grid(row=2, column=0, sticky="ew")

        tabs = ctk.CTkTabview(main)
        tabs.grid(row=3, column=0, sticky="nsew", pady=(8, 0))
        log_tab, prev_tab = tabs.add("Log de status"), tabs.add("Prévia das vagas")

        self.log_box = ctk.CTkTextbox(log_tab, state="disabled", wrap="word",
                                      font=ctk.CTkFont(family="Consolas", size=12))
        self.log_box.pack(fill="both", expand=True)
        for tag, color in (("info", "#9ca3af"), ("ok", "#22c55e"),
                           ("warn", "#f59e0b"), ("error", "#ef4444")):
            self.log_box.tag_config(tag, foreground=color)

        self.counter_label = ctk.CTkLabel(prev_tab, text="0 vagas válidas", anchor="w")
        self.counter_label.pack(fill="x")
        wrap = tk.Frame(prev_tab)
        wrap.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(wrap, columns=[c[0] for c in self.TREE_COLS],
                                 show="headings")
        for key, label, width in self.TREE_COLS:
            self.tree.heading(key, text=label)
            self.tree.column(key, width=width, anchor="w")
        vs = ttk.Scrollbar(wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vs.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vs.pack(side="right", fill="y")

    # ---- aparência --------------------------------------------------------
    def _set_appearance(self, choice: str):
        ctk.set_appearance_mode({"Escuro": "dark", "Claro": "light",
                                 "Sistema": "system"}[choice])
        self.after(150, self._apply_tree_style)

    def _apply_tree_style(self):
        dark = ctk.get_appearance_mode() == "Dark"
        bg, fg = ("#2b2b2b", "#f3f4f6") if dark else ("#ffffff", "#111827")
        st = ttk.Style(self)
        st.theme_use("clam")
        st.configure("Treeview", background=bg, fieldbackground=bg, foreground=fg,
                     rowheight=26, borderwidth=0, font=("Segoe UI", 10))
        st.configure("Treeview.Heading", background=FIAP_PINK, foreground="white",
                     font=("Segoe UI", 10, "bold"), relief="flat")
        st.map("Treeview", background=[("selected", FIAP_PINK)])

    # ---- log / eventos ----------------------------------------------------
    def _log(self, msg: str, level: str = "info"):
        stamp = datetime.now().strftime("%H:%M:%S")
        self.log_box.configure(state="normal")
        self.log_box.insert("end", f"[{stamp}] {msg}\n", level)
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    def _poll(self):
        try:
            while True:
                kind, *payload = self.events.get_nowait()
                if kind == "log":
                    self._log(*payload)
                elif kind == "progress":
                    done, total, label = payload
                    frac = min(1.0, done / total) if total else 0
                    self.progress.set(frac)
                    self.pct_label.configure(text=f"{int(frac * 100)}%")
                    self.status_label.configure(text=label)
                elif kind == "job":
                    row = payload[0]
                    self.rows.append(row)
                    self.tree.insert("", "end", values=(
                        row["Título da Oportunidade"], row["Modalidade de Contratação"],
                        row["Modalidade de Trabalho"], row["Cluster e Blocos"],
                        row["Nível da Vaga"], row["Local da Oportunidade"]))
                    self.counter_label.configure(text=f"{len(self.rows)} vagas válidas")
                elif kind == "done":
                    self._finish(*payload)
        except queue.Empty:
            pass
        self.after(100, self._poll)

    # ---- ações ------------------------------------------------------------
    def start_search(self):
        terms = [t.strip() for t in self.terms_box.get("1.0", "end").splitlines() if t.strip()]
        if not (self.var_gupy.get() or self.var_github.get() or
                self.var_remotive.get() or self.var_jooble.get()):
            messagebox.showwarning("Fontes", "Selecione ao menos uma fonte.")
            return
        if not terms and (self.var_gupy.get() or self.var_jooble.get()):
            messagebox.showwarning("Termos", "Informe ao menos um termo de busca.")
            return
        if self.var_jooble.get() and not self.jooble_key.get().strip():
            messagebox.showwarning("Jooble", "Informe a chave da API do Jooble "
                                             "ou desmarque a fonte.")
            return

        self.rows.clear()
        self.tree.delete(*self.tree.get_children())
        self.counter_label.configure(text="0 vagas válidas")
        self.progress.set(0)
        self.pct_label.configure(text="0%")
        self.stop_event.clear()
        self.btn_start.configure(state="disabled")
        self.btn_stop.configure(state="normal", fg_color="#dc2626", hover_color="#b91c1c")
        self.btn_export.configure(state="disabled")

        options = dict(
            terms=terms, pages=int(self.pages_slider.get()),
            gupy=self.var_gupy.get(), github=self.var_github.get(),
            remotive=self.var_remotive.get(), jooble=self.var_jooble.get(),
            jooble_key=self.jooble_key.get(), gh_token=self.gh_token.get())
        self.worker = threading.Thread(target=self._worker, args=(options,), daemon=True)
        self.worker.start()

    def stop_search(self):
        self.stop_event.set()
        self.btn_stop.configure(state="disabled")
        self._log("Parada solicitada… finalizando requisição atual.", "warn")

    def _worker(self, o: dict):
        put = self.events.put
        log = lambda msg, lvl="info": put(("log", msg, lvl))
        try:
            session = build_session()
            sources: list[JobSource] = []
            if o["gupy"]:
                sources.append(GupySource(session, log))
            if o["github"]:
                sources.append(GitHubIssuesSource(session, log, o["gh_token"]))
            if o["remotive"]:
                sources.append(RemotiveSource(session, log))
            if o["jooble"]:
                sources.append(JoobleSource(session, log, o["jooble_key"]))
            hunter = JobHunter(
                sources, o["terms"], o["pages"], self.stop_event, log,
                lambda d, t, l: put(("progress", d, t, l)),
                lambda row: put(("job", row)))
            # a UI acumula as linhas via evento "job"; aqui evitamos duplicar
            hunter.rows = []
            hunter.run()
            put(("done", len(hunter.seen), hunter.discarded, self.stop_event.is_set()))
        except Exception as exc:  # última barreira: o app nunca cai
            log(f"Erro inesperado na coleta: {exc!r}", "error")
            put(("done", 0, 0, True))

    def _finish(self, seen: int, discarded: int, stopped: bool):
        self.btn_start.configure(state="normal")
        self.btn_stop.configure(state="disabled", fg_color="#6b7280", hover_color="#4b5563")
        if self.rows:
            self.btn_export.configure(state="normal")
        if not stopped:
            self.progress.set(1)
            self.pct_label.configure(text="100%")
        self.status_label.configure(text="Interrompido." if stopped else "Concluído.")
        self._log(f"{'Interrompido' if stopped else 'Concluído'}: {len(self.rows)} vagas "
                  f"válidas · {seen} únicas analisadas · {discarded} descartadas "
                  f"(fora de tecnologia).", "ok")

    def export_excel(self):
        if not self.rows:
            messagebox.showinfo("Exportar", "Não há vagas para exportar.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"vagas_fiap_{date.today():%Y%m%d}.xlsx")
        if not path:
            return
        try:
            export_to_excel(self.rows, path)
        except PermissionError:
            messagebox.showerror("Exportar", "Não foi possível gravar o arquivo. "
                                             "Ele está aberto no Excel?")
            self._log("Falha ao exportar: arquivo em uso.", "error")
        except Exception as exc:
            messagebox.showerror("Exportar", f"Erro ao exportar: {exc}")
            self._log(f"Falha ao exportar: {exc!r}", "error")
        else:
            self._log(f"Excel gerado com {len(self.rows)} vagas: {path}", "ok")
            messagebox.showinfo("Exportar", f"Arquivo salvo com {len(self.rows)} vagas.")

    def _on_close(self):
        self.stop_event.set()
        self.destroy()


# --------------------------------------------------------------------------- #
# Instância de teste (offline)
# --------------------------------------------------------------------------- #
def run_selftest(out_path: str = "vagas_fiap_selftest.xlsx") -> int:
    today = date.today().isoformat()
    samples = [
        RawJob("Teste", "Analista de Segurança da Informação Sênior (SOC) – Vaga afirmativa "
               "para Pessoas Pretas e Pardas", "https://exemplo.com/vaga/1",
               "<p>Atuação em <b>SIEM</b> e resposta a incidentes. Vaga afirmativa para "
               "pessoas pretas e pardas. Modelo híbrido, 40 horas semanais. "
               "Salário: R$ 9.000 a R$ 12.000.</p>", "Banco Exemplo", "São Paulo - SP",
               today, None, "vacancy_type_effective", "hybrid"),
        RawJob("Teste", "Estagiário de Dados", "https://exemplo.com/vaga/2",
               "<ul><li>SQL</li><li>Power BI</li></ul> Bolsa R$ 2.200 por mês. 6 horas diárias.",
               "Varejo SA", "Campinas - SP", today, None, "vacancy_type_internship", "on-site"),
        RawJob("Teste", "[Remoto] Desenvolvedor Backend Python Pleno", "https://exemplo.com/vaga/3",
               "Contratação PJ. Python, Django, AWS. 100% remoto.", "Startup X", "",
               "2025-01-10", "2025-02-01", None, None, "PJ Remoto"),
        RawJob("Teste", "Gerente de Engenharia de Software", "https://exemplo.com/vaga/4",
               "Liderar times de desenvolvimento backend e frontend. Buscamos pessoas com "
               "deficiência (PCD) - vagas inclusivas. Presencial.", "Fintech Y",
               "Rio de Janeiro - RJ", today, None, "vacancy_type_effective"),
        RawJob("Teste", "Assistente Administrativo", "https://exemplo.com/vaga/5",
               "Rotinas de escritório e atendimento.", "Empresa Z", "Curitiba - PR"),
        RawJob("Teste", "Engenheiro de IA / Machine Learning " + "Especialista " * 6,
               "https://exemplo.com/vaga/6", "LLM, PyTorch, MLOps. " * 200,
               "Lab AI", "Remoto", today, None, None, "remote"),
    ]
    rows = []
    for s in samples:
        row = process_job(s)
        if row is None:
            print(f"- descartada (não-tech): {s.title}")
            continue
        issues = validate_row(row)
        print(f"+ {row['Título da Oportunidade']!r}\n    {row['Modalidade de Contratação']} | "
              f"{row['Modalidade de Trabalho']} | {row['Cluster e Blocos']} | "
              f"{row['Nível da Vaga']} | {row['Data de Expiração']:%d/%m/%Y} | "
              f"título={len(row['Título da Oportunidade'])} desc={len(row['Descrição da Oportunidade'])}"
              f" | problemas={issues or 'nenhum'}")
        assert not issues, issues
        rows.append(row)
    export_to_excel(rows, out_path)
    print(f"\nExcel de teste gerado: {out_path} ({len(rows)} linhas)")
    return 0


def main() -> int:
    if "--selftest" in sys.argv:
        idx = sys.argv.index("--selftest")
        out = sys.argv[idx + 1] if len(sys.argv) > idx + 1 else "vagas_fiap_selftest.xlsx"
        return run_selftest(out)
    if ctk is None:
        print("CustomTkinter não instalado. Execute: pip install -r requirements.txt")
        return 1
    App().mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())